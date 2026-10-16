#!/usr/bin/env python3
"""Update one row in the Project Tracker (#TRKR) .xlsx without loading it into an agent's context.

Usage (paths may be relative or absolute; use / separators):
  py tools/trkr_update.py --tracker "C:/.../Project-Tracker-SC.xlsx" --sheet WS-1.1 \
      --id PKG-1.1.3 --status Complete --note "minor spillover: X"

- Column headers are matched case-insensitively; override with --id-col/--status-col/--completed-col/--notes-col.
- Without --sheet every tab is searched (IDs are unique project-wide).
- A lock file (<tracker>.lock) serialises parallel Team Leads; stale locks (>120 s) are cleared.
- A .bak copy is written before every save (openpyxl drops charts/images; none are expected).
Exit codes: 0 ok | 2 row/column/sheet not found | 3 file locked (Excel or another lead) | 4 other error.
Prints exactly one line. Never prints sheet contents.
"""
import argparse, os, shutil, sys, time
from datetime import datetime

try:
    from openpyxl import load_workbook
except ImportError:
    print("ERROR openpyxl not installed (py -m pip install openpyxl)"); sys.exit(4)

LOCK_WAIT_S, LOCK_STALE_S = 60, 120

def acquire_lock(path):
    lock = path + ".lock"
    deadline = time.time() + LOCK_WAIT_S
    while True:
        try:
            fd = os.open(lock, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
            os.write(fd, f"{os.getpid()} {datetime.now().isoformat()}".encode())
            os.close(fd)
            return lock
        except FileExistsError:
            try:
                if time.time() - os.path.getmtime(lock) > LOCK_STALE_S:
                    os.remove(lock); continue
            except FileNotFoundError:
                continue
            if time.time() > deadline:
                print(f"LOCKED {path} (another Team Lead is updating; retry)"); sys.exit(3)
            time.sleep(1)

def find_header(ws, id_col, max_scan=15):
    for r in range(1, min(ws.max_row, max_scan) + 1):
        heads = {str(c.value).strip().lower(): c.column for c in ws[r] if c.value is not None}
        if id_col.lower() in heads:
            return r, heads
    return None, None

def run(a):
    for attempt in (1, 2):
        try:
            wb = load_workbook(a.tracker); break
        except PermissionError:
            if attempt == 2:
                print(f"LOCKED {a.tracker} (close it in Excel and retry)"); sys.exit(3)
            time.sleep(5)
        except FileNotFoundError:
            print(f"ERROR tracker not found: {a.tracker}"); sys.exit(4)

    if a.sheet and a.sheet not in wb.sheetnames:
        print(f"NOTFOUND sheet '{a.sheet}'"); sys.exit(2)
    sheets = [wb[a.sheet]] if a.sheet else wb.worksheets
    target = None
    for ws in sheets:
        hr, heads = find_header(ws, a.id_col)
        if hr is None:
            continue
        idc = heads[a.id_col.lower()]
        for r in range(hr + 1, ws.max_row + 1):
            v = ws.cell(r, idc).value
            if v is not None and str(v).strip() == a.id:
                target = (ws, r, heads); break
        if target is None and a.create:
            r = ws.max_row + 1
            ws.cell(r, idc, a.id)
            target = (ws, r, heads)
        if target:
            break
    if target is None:
        print(f"NOTFOUND id={a.id} (id column '{a.id_col}')"); sys.exit(2)

    ws, r, heads = target
    sc = heads.get(a.status_col.lower())
    if sc is None:
        print(f"NOTFOUND column '{a.status_col}' on sheet '{ws.title}'"); sys.exit(2)
    ws.cell(r, sc, a.status)
    cc = heads.get(a.completed_col.lower())
    if cc and a.status.lower() == "complete":
        ws.cell(r, cc, datetime.now().replace(microsecond=0))
    if a.note:
        nc = heads.get(a.notes_col.lower())
        if nc is None:
            print(f"NOTFOUND column '{a.notes_col}' on sheet '{ws.title}'"); sys.exit(2)
        old = ws.cell(r, nc).value
        ws.cell(r, nc, f"{old}; {a.note}" if old else a.note)

    if a.dry_run:
        print(f"DRYRUN {a.id} -> {a.status} sheet='{ws.title}' row={r}"); return
    for attempt in (1, 2):
        try:
            if not a.no_backup:
                shutil.copy2(a.tracker, a.tracker + ".bak")
            wb.save(a.tracker); break
        except PermissionError:
            if attempt == 2:
                print(f"LOCKED {a.tracker} (close it in Excel and retry)"); sys.exit(3)
            time.sleep(5)
    print(f"OK {a.id} -> {a.status} sheet='{ws.title}' row={r}")

def main():
    p = argparse.ArgumentParser()
    p.add_argument("--tracker", required=True)
    p.add_argument("--id", required=True, help="Row identifier, e.g. PKG-1.1.3 or WS-1.1")
    p.add_argument("--status", required=True)
    p.add_argument("--note", default="")
    p.add_argument("--sheet", default=None, help="Tab name (e.g. WS-1.1); default searches all tabs")
    p.add_argument("--id-col", default="ID")
    p.add_argument("--status-col", default="Status")
    p.add_argument("--completed-col", default="Completed")
    p.add_argument("--notes-col", default="Notes")
    p.add_argument("--create", action="store_true", help="Append a row if the ID is not found")
    p.add_argument("--no-backup", action="store_true")
    p.add_argument("--dry-run", action="store_true")
    a = p.parse_args()
    lock = None if a.dry_run else acquire_lock(a.tracker)
    try:
        run(a)
    finally:
        if lock and os.path.exists(lock):
            os.remove(lock)

if __name__ == "__main__":
    try:
        main()
    except SystemExit:
        raise
    except Exception as e:
        print(f"ERROR {type(e).__name__}: {e}"); sys.exit(4)
